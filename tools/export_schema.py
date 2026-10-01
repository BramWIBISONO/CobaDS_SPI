"""Ekspor definisi tabel Excel v4 ke importer/schema/excel_tables.json (dijalankan sekali, hanya membaca).

Tata letak tabel = v2_layout.py pipeline (judul kolom, kunci, jenis kolom, format angka). Jenis nilai (text / number / date / datetime /
time / bool) ditentukan dari format; bila kolom tidak berformat tanggal/angka, dari nilai di workbook Jakarta. Kolom yang isinya campur
(angka/tanggal + teks) mendapat kolom pendamping <nama>_text. Sheet staging D_BULAN / D_MURID (judul di baris 3) dibaca dari workbook;
kolom rumusnya dikenali dari selnya.
    .venv/Scripts/python tools/export_schema.py
"""
import datetime
import json
import keyword
import re
import sys
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl

WEB = Path(__file__).resolve().parent.parent
MGMT = WEB.parent
PIPE = MGMT / "00_SYSTEM" / "APP_BUILD_V2" / "_scripts"
JKT = MGMT / "APP" / "SPI_STUDENT-SPP_APP_2026_09_v4.xlsm"
OUT = WEB / "importer" / "schema" / "excel_tables.json"
sys.path.insert(0, str(PIPE))
import v2_layout as LY  # noqa: E402

APP_OF = {
    "PROGRAM_MASTER": "masterdata", "OFF_REASON_MASTER": "masterdata", "TARIF_FEE": "masterdata", "UNIT_MASTER": "masterdata",
    "TEACHER_MASTER": "masterdata", "ROOM_MASTER": "masterdata", "PARTNER_MASTER": "masterdata",
    "STUDENT_MASTER": "students", "PARENT_MASTER": "students", "STATUS_EVENT": "students", "STUDENT_OFF": "students",
    "FOLLOW_UP": "students", "ACADEMIC_RECORD": "students", "DOKUMEN": "students", "LEAD": "students",
    "STUDENT_ID_MAPPING": "students", "D_BULAN": "students", "D_MURID": "students",
    "CLASS_MASTER": "classes", "CLASS_MEMBERS": "classes", "CLASS_SCHEDULE": "classes", "SESI": "classes", "SIMULASI_JADWAL": "classes",
    "PERIODE": "finance", "SPP_TAGIHAN": "finance", "BUKTI_BAYAR": "finance", "BUKU_KAS": "finance", "BUKU_KAS_KELUAR": "finance",
    "PEMBAYAR": "finance", "NOTA_LOG": "finance",
    "ISSUE_UNIT": "quality",
    "AUDIT_LOG": "audit", "IMPORT_LOG": "audit", "SOURCE_REFERENCE": "audit",
}
STAGING = {
    "D_BULAN": {"header_row": 3, "key_header": "Kunci (ID|yyyymm)", "aliases": {},
                "names": {"Kunci (ID|yyyymm)": "key", "ID": "v1", "Nama Murid": "nama", "Bulan": "bulan", "Status": "status",
                          "Tanggal status": "tgl_status", "Keterangan": "ket", "Level (Grade)": "grade", "Program": "program",
                          "Guru (asli)": "guru_asli", "Guru (rapi)": "guru", "Tipe Kelas": "tipe", "Mode": "mode", "Kelompok": "kelompok",
                          "Sumber (sel asli)": "sumber", "Student ID": "std", "Kode Kelas (bulan itu)": "kode_bulan",
                          "Kode Kelas (terakhir, DB Murid)": "kode_terakhir", "Sumber Kode Kelas": "kode_sumber", "Sekolah": "sekolah",
                          "Issue Terbuka (bulan ini)": "issues", "Import Batch": "batch"}},
    "D_MURID": {"header_row": 3, "key_header": "ID", "names": {"ID": "v1", "Student ID": "std", "Nama Murid": "nama"},
                "aliases": {"Baris di DB Murid JKT": ["Baris di DB Murid cabang"]}},
}
NON_UNIQUE = {"IMPORT_LOG"}                       # satu batch = beberapa baris (satu per file)
SKIP_KEY_ONLY = {"SIMULASI_JADWAL"}               # baris bernomor tanpa isi
STOP_AT_BLANK = {"SIMULASI_JADWAL"}               # blok 20 baris simulasi; di bawahnya tampilan JADWAL RESMI (bukan data)
KIND_OVERRIDES = {("SESI", "hadir"): "number"}    # tabel kosong di Jakarta: jenis tidak bisa dibaca dari data
for _k in ("lid", "user", "action", "entity", "eid", "field", "old", "new", "by"):
    KIND_OVERRIDES[("AUDIT_LOG", _k)] = "text"
for _k in ("batch", "date", "file", "path", "sha", "sheets", "rows", "excluded", "ver", "notes"):
    KIND_OVERRIDES[("IMPORT_LOG", _k)] = "text"
RESERVED = {"id", "pk", "branch", "row_no", "objects", "save", "delete", "clean", "check", "full_clean"}


def model_name(sheet):
    return {"D_BULAN": "DBulan", "D_MURID": "DMurid"}.get(sheet) or "".join(p.capitalize() for p in sheet.lower().split("_"))


def safe(name):
    name = re.sub(r"\W+", "_", name).strip("_").lower() or "kolom"
    if name[0].isdigit():
        name = "k_" + name
    if keyword.iskeyword(name) or name in RESERVED:
        name += "_x"
    return name


def kind_from_fmt(fmt):
    if not fmt:
        return None
    f = fmt.lower()
    if "hh" in f and ("dd" in f or "yy" in f):
        return "datetime"
    if "hh" in f:
        return "time"
    if "dd" in f or "yy" in f or "mmm" in f:
        return "date"
    if "0" in f or "#" in f:
        return "number"
    return None


def tag(v):
    if isinstance(v, bool):
        return "bool"
    if isinstance(v, (int, float)):
        return "number"
    if isinstance(v, datetime.datetime):
        return "datetime" if (v.hour, v.minute, v.second) != (0, 0, 0) else "date"
    if isinstance(v, datetime.date):
        return "date"
    if isinstance(v, datetime.time):
        return "time"
    return "text"


def decide(fmt, found):
    kind = kind_from_fmt(fmt)
    if kind is None:
        kind = next((k for k in ("datetime", "date", "time", "number", "bool") if k in found), "text")
    if kind == "date" and "datetime" in found:
        kind = "datetime"
    return kind, (kind != "text" and "text" in found)


def sample(ws, header_row, key_idx, stop_at_blank=False):
    out = defaultdict(Counter)
    for row in ws.iter_rows(min_row=header_row + 1, values_only=True):
        if key_idx >= len(row) or row[key_idx] in (None, ""):
            if stop_at_blank:
                break                                 # sama dengan reader: blok berakhir di baris kosong pertama
            continue
        for i, v in enumerate(row):
            if v not in (None, ""):
                out[i][tag(v)] += 1
    return out


def header_list(ws, row):
    cells = next(ws.iter_rows(min_row=row, max_row=row, values_only=True), ())
    return [str(h).strip() if h is not None else "" for h in cells]


def layout_tables(wb):
    out = []
    for sheet, t in LY.TABLES.items():
        if sheet == "SETTINGS":
            continue                                  # branches.BranchSetting + identitas Branch
        ws = wb[sheet]
        hdr = header_list(ws, LY.HDR)
        tags = sample(ws, LY.HDR, hdr.index(t["key"]), sheet in STOP_AT_BLANK)
        fields, names, key_name = [], set(), None
        for header, key, avail, _d, _s, _w, fmt, _h in t["cols"]:
            name = safe(key)
            assert name not in names, (sheet, name)
            names.add(name)
            i = hdr.index(header) if header in hdr else None
            kind, mixed = decide(fmt, set(tags.get(i, {})) if i is not None else set())
            kind = KIND_OVERRIDES.get((sheet, key), kind)
            mixed = mixed and kind != "text"
            stored = avail != "FORMULA"
            fields.append({"header": header, "name": name, "avail": avail, "kind": kind, "stored": stored,
                           "text_field": f"{name}_text" if (stored and mixed) else None, "aliases": []})
            if header == t["key"]:
                key_name = name
        out.append({"sheet": sheet, "app": APP_OF[sheet], "model": model_name(sheet), "header_row": LY.HDR, "data_row": LY.D0,
                    "key": key_name, "unique": sheet not in NON_UNIQUE, "skip_key_only": sheet in SKIP_KEY_ONLY,
                    "stop_at_blank": sheet in STOP_AT_BLANK, "fields": fields})
    return out


def staging_tables(wb, wbf):
    out = []
    for sheet, cfg in STAGING.items():
        hr = cfg["header_row"]
        hdr = header_list(wb[sheet], hr)
        hdr = hdr[:max(i for i, h in enumerate(hdr) if h) + 1]
        tags = sample(wb[sheet], hr, hdr.index(cfg["key_header"]))
        formula_cols = set()
        for row in wbf[sheet].iter_rows(min_row=hr + 1, max_row=hr + 200, values_only=True):
            formula_cols |= {i for i, v in enumerate(row[:len(hdr)]) if isinstance(v, str) and v.startswith("=")}
        fields, names = [], set()
        for i, header in enumerate(hdr):
            if not header:
                continue
            name = base = safe(cfg["names"].get(header, header))
            n = 2
            while name in names:
                name, n = f"{base}_{n}", n + 1
            names.add(name)
            stored = i not in formula_cols
            kind, mixed = decide(None, set(tags.get(i, {})))
            fields.append({"header": header, "name": name, "avail": "ASLI" if stored else "FORMULA", "kind": kind, "stored": stored,
                           "text_field": f"{name}_text" if (stored and mixed) else None, "aliases": cfg["aliases"].get(header, [])})
        out.append({"sheet": sheet, "app": APP_OF[sheet], "model": model_name(sheet), "header_row": hr, "data_row": hr + 1,
                    "key": safe(cfg["names"].get(cfg["key_header"], cfg["key_header"])), "unique": True, "skip_key_only": False,
                    "stop_at_blank": False,
                    "fields": fields})
    return out


def main():
    wb = openpyxl.load_workbook(JKT, read_only=True, data_only=True)
    wbf = openpyxl.load_workbook(JKT, read_only=True, data_only=False)
    tables = layout_tables(wb) + staging_tables(wb, wbf)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps({"version": 1, "source": {"layout": "00_SYSTEM/APP_BUILD_V2/_scripts/v2_layout.py", "workbook": JKT.name},
                               "tables": tables}, ensure_ascii=False, indent=1), encoding="utf-8")
    print("tabel", len(tables), "| kolom tersimpan", sum(1 for t in tables for f in t["fields"] if f["stored"]),
          "| kolom campur", [(t["sheet"], f["header"]) for t in tables for f in t["fields"] if f["text_field"]])


if __name__ == "__main__":
    main()
