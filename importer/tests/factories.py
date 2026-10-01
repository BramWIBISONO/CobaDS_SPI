"""Workbook SPI v4 minimal untuk uji: setiap sheet skema dengan baris judulnya; isi = {sheet: [{judul kolom: nilai}]}."""
import openpyxl

from importer.reader import SETTINGS_HEADER_ROW, SETTINGS_SHEET
from importer.schema import load_schema


def build_workbook(path, rows=None, settings=None, drop_sheets=(), rename_headers=None):
    rows, rename_headers = rows or {}, rename_headers or {}
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for spec in load_schema():
        if spec.sheet in drop_sheets:
            continue
        ws = wb.create_sheet(spec.sheet)
        for j, f in enumerate(spec.fields, 1):
            ws.cell(row=spec.header_row, column=j, value=rename_headers.get((spec.sheet, f.header), f.header))
        for i, record in enumerate(rows.get(spec.sheet, [])):
            for j, f in enumerate(spec.fields, 1):
                if f.header in record:
                    ws.cell(row=spec.data_row + i, column=j, value=record[f.header])
    ws = wb.create_sheet(SETTINGS_SHEET)
    for j, h in enumerate(("Parameter", "Nilai", "Keterangan", "Kunci"), 1):
        ws.cell(row=SETTINGS_HEADER_ROW, column=j, value=h)
    for i, (key, value) in enumerate((settings if settings is not None else {"unit": "UNIT-JKT"}).items()):
        r = SETTINGS_HEADER_ROW + 1 + i
        ws.cell(row=r, column=1, value=key)
        ws.cell(row=r, column=2, value=value)
        ws.cell(row=r, column=4, value=key)
    wb.save(path)
    return path
