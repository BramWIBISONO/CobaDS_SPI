"""Membaca workbook SPI v4 lewat judul kolom (sama dengan pipeline). Hanya kolom tersimpan yang dibaca; nilai rumus diabaikan."""
import zipfile
from dataclasses import dataclass, field

import openpyxl
from openpyxl.utils.exceptions import InvalidFileException

from .schema import load_schema

SETTINGS_SHEET = "SETTINGS"
SETTINGS_HEADER_ROW = 5


class WorkbookError(Exception):
    pass


@dataclass
class SheetRow:
    row: int
    values: dict


@dataclass
class WorkbookData:
    tables: dict = field(default_factory=dict)
    missing_headers: dict = field(default_factory=dict)
    settings: list = field(default_factory=list)

    def setting(self, key):
        return next((s for s in self.settings if s["key"] == key), None)

    @property
    def unit(self):
        s = self.setting("unit")
        return str(s["value"]).strip() if s and s["value"] not in (None, "") else None


def _open(path, data_only):
    try:
        return openpyxl.load_workbook(path, read_only=True, data_only=data_only)
    except (InvalidFileException, zipfile.BadZipFile, KeyError, OSError, ValueError) as exc:
        raise WorkbookError("File bukan workbook Excel (.xlsx / .xlsm) yang bisa dibaca.") from exc


def _headers(ws, row):
    cells = next(ws.iter_rows(min_row=row, max_row=row, values_only=True), ())
    return {str(h).strip(): i for i, h in enumerate(cells) if h not in (None, "")}


def _read_table(ws, spec, data):
    hdr = _headers(ws, spec.header_row)
    cols, missing = {}, []
    for f in spec.stored_fields:
        i = hdr.get(f.header)
        if i is None:
            i = next((hdr[a] for a in f.aliases if a in hdr), None)
        if i is None:
            missing.append(f.header)
        else:
            cols[f.name] = i
    if missing:
        data.missing_headers[spec.sheet] = missing
    rows = []
    key_i = cols.get(spec.key)
    if key_i is not None:
        for r, cells in enumerate(ws.iter_rows(min_row=spec.data_row, values_only=True), spec.data_row):
            key = cells[key_i] if key_i < len(cells) else None
            if key in (None, ""):
                if spec.stop_at_blank:
                    break                          # akhir blok (SIMULASI_JADWAL: di bawahnya tampilan JADWAL RESMI, bukan data)
                continue
            values = {name: (cells[i] if i < len(cells) else None) for name, i in cols.items()}
            if spec.skip_key_only and all(v in (None, "") for n, v in values.items() if n != spec.key):
                continue
            rows.append(SheetRow(r, values))
    data.tables[spec.sheet] = rows


def _read_settings(ws_values, ws_formulas, data):
    hdr = _headers(ws_values, SETTINGS_HEADER_ROW)
    need = ("Parameter", "Nilai", "Keterangan", "Kunci")
    if any(h not in hdr for h in need):
        data.missing_headers[SETTINGS_SHEET] = [h for h in need if h not in hdr]
        return
    formula_rows = set()
    for r, cells in enumerate(ws_formulas.iter_rows(min_row=SETTINGS_HEADER_ROW + 1, values_only=True), SETTINGS_HEADER_ROW + 1):
        v = cells[hdr["Nilai"]] if hdr["Nilai"] < len(cells) else None
        if isinstance(v, str) and v.startswith("="):
            formula_rows.add(r)
    for r, cells in enumerate(ws_values.iter_rows(min_row=SETTINGS_HEADER_ROW + 1, values_only=True), SETTINGS_HEADER_ROW + 1):
        get = lambda h: cells[hdr[h]] if hdr[h] < len(cells) else None   # noqa: E731
        if get("Kunci") in (None, ""):
            continue
        data.settings.append({"key": str(get("Kunci")), "param": get("Parameter") or "", "value": get("Nilai"),
                              "note": get("Keterangan") or "", "row": r, "is_formula": r in formula_rows})


def read_workbook(path):
    wb = _open(path, data_only=True)
    try:
        names = set(wb.sheetnames)
        missing = [s for s in [t.sheet for t in load_schema()] + [SETTINGS_SHEET] if s not in names]
        if missing:
            shown = ", ".join(missing[:5]) + (f" dan {len(missing) - 5} lainnya" if len(missing) > 5 else "")
            raise WorkbookError(f"Bukan workbook SPI v4: sheet {shown} tidak ada.")
        data = WorkbookData()
        for spec in load_schema():
            _read_table(wb[spec.sheet], spec, data)
        wbf = _open(path, data_only=False)
        try:
            _read_settings(wb[SETTINGS_SHEET], wbf[SETTINGS_SHEET], data)
        finally:
            wbf.close()
        return data
    finally:
        wb.close()
